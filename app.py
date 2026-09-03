from flask import Flask, render_template, request, Response
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

ITEMS_POR_MAQUINA = [
    {"id": "enciende", "label": "Enciende correctamente"},
    {"id": "completa", "label": "Está completa (Mouse, Teclado, Monitor)"},
    {"id": "red", "label": "Cuenta con conexión a Internet"}
]

@app.route('/')
def index():
    return render_template('index.html', items_maquina=ITEMS_POR_MAQUINA)

@app.route('/exportar', methods=['POST'])
def exportar():
    data = request.form
    agente = data.get('agente', 'No especificado').strip()
    letra_sala = data.get('letra_sala', 'A').strip().upper()
    num_sala = data.get('num_sala', '101').strip()
    num_maquinas = int(data.get('num_maquinas', 0))
    observaciones = data.get('observaciones', '').strip()
    
    fecha_actual = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    codigo_sala = f"Sala {letra_sala}-{num_sala}"

    # Crear libro Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Revisión"
    ws.views.sheetView[0].showGridLines = True

    # Estilos
    font_titulo = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10)
    
    fill_titulo = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    fill_ok = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")   # Verde claro
    fill_warn = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Amarillo claro
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Encabezado principal
    ws.merge_cells("A1:G1")
    cell_title = ws["A1"]
    cell_title.value = f"REPORTE DE CHECKLIST - {codigo_sala.upper()}"
    cell_title.font = font_titulo
    cell_title.fill = fill_titulo
    cell_title.alignment = align_center
    ws.row_dimensions[1].height = 35

    # 2. Metadata de la sala y agente
    ws["A3"] = "Agente / Técnico:"
    ws["B3"] = agente
    ws["A4"] = "Edificio / Letra:"
    ws["B4"] = letra_sala
    ws["A5"] = "Número de Sala:"
    ws["B5"] = num_sala

    ws["E3"] = "Fecha de Revisión:"
    ws["F3"] = fecha_actual
    ws["E4"] = "Total Equipos:"
    ws["F4"] = num_maquinas

    for cell in ["A3", "A4", "A5", "E3", "E4"]:
        ws[cell].font = font_bold
    for cell in ["B3", "B4", "B5", "F3", "F4"]:
        ws[cell].font = font_regular

    # 3. Encabezados de la tabla
    headers = ["Fecha y Hora", "Agente", "Equipo", "Enciende", "Esta Completa", "Tiene Internet", "Estado PC"]
    start_row = 7
    
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 24

    # 4. Filas por máquina
    current_row = start_row + 1
    for i in range(1, num_maquinas + 1):
        enciende = "SI" if f"pc_{i}_enciende" in data else "NO"
        completa = "SI" if f"pc_{i}_completa" in data else "NO"
        red = "SI" if f"pc_{i}_red" in data else "NO"
        
        es_ok = (enciende == "SI" and completa == "SI" and red == "SI")
        estado_pc = "OK" if es_ok else "CON REVISIÓN PENDIENTE"

        row_values = [fecha_actual, agente, f"PC {i:02d}", enciende, completa, red, estado_pc]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            cell.alignment = align_center if col_idx not in [1, 2] else align_left
            
            if col_idx == 7:
                cell.fill = fill_ok if es_ok else fill_warn
                cell.font = font_bold

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # 5. Observaciones
    current_row += 1
    ws.cell(row=current_row, column=1, value="OBSERVACIONES GENERALES:").font = font_bold
    current_row += 1
    
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row + 2, end_column=7)
    obs_cell = ws.cell(row=current_row, column=1, value=observaciones if observaciones else "Sin observaciones registradas.")
    obs_cell.font = font_regular
    obs_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Ajuste automático de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws.column_dimensions['G'].width = 24

    # Generar respuesta
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Checklist_Sala_{letra_sala}-{num_sala}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)